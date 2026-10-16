#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script de configuración y pruebas.
Crea carpetas, crea Datos_informe_veterinario.xlsx con la hoja 'Datos' (plantilla),
crea Plantilla_informe_veterinario.docx de ejemplo, genera imágenes de prueba y
puedes ejecutar generar_informe.py para validar el flujo.
"""

import os
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from docx import Document
from docx.shared import Pt
from PIL import Image, ImageDraw, ImageFont
import subprocess
import sys

PROJECT_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH = os.path.join(PROJECT_DIR, "Datos_informe_veterinario.xlsx")
TEMPLATE_PATH = os.path.join(PROJECT_DIR, "Plantilla_informe_veterinario.docx")
FOTOGRAFIAS_DIR = os.path.join(PROJECT_DIR, "Fotografias")
IMG_SUBDIRS = ["Examen_externo", "Examen_interno", "Examen_radiografico"]
GENERAR_SCRIPT = os.path.join(PROJECT_DIR, "generar_informe.py")

def create_dirs():
    os.makedirs(FOTOGRAFIAS_DIR, exist_ok=True)
    for d in IMG_SUBDIRS:
        os.makedirs(os.path.join(FOTOGRAFIAS_DIR, d), exist_ok=True)
    os.makedirs(os.path.join(PROJECT_DIR, "Informes_generados"), exist_ok=True)

def create_excel_template(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"
    # Fill column A with labels and column B with empty cells or example values
    labels = [
        "Número de caso","Número de historia","Especie","Edad","Sexo","Identificación del animal",
        "Paraje","Término municipal","Provincia","UTV","Coordenadas","Número de precinto o acta",
        "Fecha de ingreso en el CRFS AMUS","Tipo de ingreso","Fecha de necropsia","Conservación del cadáver",
        "Examen radiográfico realizado","Análisis toxicológico realizado",
        "Examen externo","Examen interno","Resultado del examen radiográfico",
        "Laboratorio de toxicología","Fecha del análisis toxicológico","Muestras enviadas",
        "Resultados toxicológicos","Diagnóstico y conclusiones",
        "Fecha de emisión","Nombre del veterinario firmante","Número de colegiado",
        "Provincia del colegio profesional","Centro de trabajo"
    ]
    for i, lab in enumerate(labels, start=1):
        ws[f"A{i}"] = lab
        # example for some fields
        if lab == "Sexo":
            ws[f"B{i}"] = "Hembra"
        if lab == "Examen radiográfico realizado":
            ws[f"B{i}"] = "No"
        if lab == "Análisis toxicológico realizado":
            ws[f"B{i}"] = "No"
    # Add a second sheet with instrucciones
    ws2 = wb.create_sheet("Instrucciones")
    ws2["A1"] = "Instrucciones para rellenar el Excel:"
    ws2["A2"] = "- Complete los campos en la columna B."
    ws2["A3"] = "- Los campos 'Sexo' y 'Examen radiográfico realizado' tienen ejemplos."
    # Save
    wb.save(path)
    print("Se ha creado el Excel plantilla en:", path)

def create_word_template(path):
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    doc.add_heading("Plantilla informe veterinario (ejemplo)", level=1)
    doc.add_paragraph("Este documento es una plantilla de ejemplo creada por setup_project.py.")
    doc.add_paragraph("Reemplace este archivo por su plantilla institucional si la tiene. El script respetará la estructura general.")
    doc.add_paragraph("\n\n\nFirma:\n\n______________________________\nNombre: \nColegiado:")
    doc.save(path)
    print("Plantilla Word de ejemplo creada en:", path)

def create_sample_images():
    # Create small images in each folder with different formats
    for idx, sub in enumerate(IMG_SUBDIRS, start=1):
        folder = os.path.join(FOTOGRAFIAS_DIR, sub)
        # generate 6 images of different names and formats
        for n in range(1, 7):
            mode = "RGB"
            img = Image.new(mode, (1200, 800), color=(200+idx*10, 120+n*10, 150))
            d = ImageDraw.Draw(img)
            text = f"{sub} {n}"
            try:
                fnt = ImageFont.load_default()
                d.text((50,50), text, font=fnt, fill=(0,0,0))
            except:
                d.text((50,50), text, fill=(0,0,0))
            # alternate formats
            if n % 5 == 0:
                ext = "png"
            elif n % 4 == 0:
                ext = "tiff"
            elif n % 3 == 0:
                ext = "webp"
            else:
                ext = "jpg"
            filename = os.path.join(folder, f"{n}.{ext}")
            img.save(filename)
    print("Imágenes de prueba creadas en:", FOTOGRAFIAS_DIR)

def run_generator():
    # Attempt to run generar_informe.py
    print("Ejecutando generación de informe de prueba...")
    try:
        subprocess.check_call([sys.executable, GENERAR_SCRIPT])
    except subprocess.CalledProcessError as e:
        print("El script generar_informe.py devolvió un error:", e)
    except Exception as e:
        print("No se pudo ejecutar el script de generación automáticamente. Ejecútelo manualmente con:")
        print(f"python {GENERAR_SCRIPT}")

if __name__ == "__main__":
    create_dirs()
    create_excel_template(EXCEL_PATH)
    create_word_template(TEMPLATE_PATH)
    create_sample_images()
    print("Configuración completada.\nSugerencia: revise Datos_informe_veterinario.xlsx, edite el 'Número de historia' y 'Especie' para pruebas reales.")
    # No ejecutamos automáticamente la generación para permitir al usuario revisar; descomentar si se desea ejecución automática:
    # run_generator()